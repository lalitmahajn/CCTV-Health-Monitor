"""
Feed real camera configuration from SCPL_FW_NVR.xlsx into SQLite database
"""
import openpyxl
import asyncio
import re
from app.database import init_db
from app.models import CameraRepository

# NVR Name aliases to match between sheets
NVR_ALIASES = {
    'admin dvr 1': 'Admin DVR 1',
    'ethenol dvr 2': 'Ethenol DVR 2',
    'it room dvr 3': 'IT Room DVR 3',
    'maitenance dvr 5': 'Maintenance DVR 5',
    'maintenance dvr 5': 'Maintenance DVR 5',
    'sanitizer plant dvr 6': 'Sanitizer Plant DVR 6',
    'sec. cabin dvr 7': 'Security Cabin DVR 7',
    'security cabin dvr 7': 'Security Cabin DVR 7',
    'store dvr 8': 'Stores DVR 8',
    'stores dvr 8': 'Stores DVR 8',
    'polymer dvr 10': 'Polymer DVR 10',
    'new stores and windsor': 'New Stores and Windsor',
    'old ms & sa nvr': 'Old MS & SA NVR',
    'new ms & os nvr': 'New MS & OS NVR',
    'aspirin 1': 'Aspirin 1',
    'aspirin 2': 'Aspirin 2',
}

async def import_real_cameras():
    db_path = "cctv_monitor.db"
    await init_db(db_path)
    repo = CameraRepository(db_path)
    from app.models import NvrRepository
    nvr_repo = NvrRepository(db_path)
    
    # Clear current dummy cameras
    existing = await repo.get_all()
    for c in existing:
        await repo.delete(c["id"])
    print(f"Cleared {len(existing)} existing records from database.")

    wb = openpyxl.load_workbook('SCPL_FW_NVR.xlsx', data_only=True)
    ws_nvr = wb['NVR_Factory']
    ws_cam = wb['CAM_Details']

    # 1. Parse NVRs
    nvrs = {}
    for r in range(2, ws_nvr.max_row + 1):
        loc = ws_nvr.cell(r, 3).value
        local_ip = ws_nvr.cell(r, 7).value
        rtsp_port = ws_nvr.cell(r, 12).value
        total_ch = ws_nvr.cell(r, 8).value
        used_ch = ws_nvr.cell(r, 9).value
        make = ws_nvr.cell(r, 4).value
        model = ws_nvr.cell(r, 5).value
        
        if not loc:
            continue
        loc_str = str(loc).strip()
        
        # Skip Sankeshwar Godowns as requested
        if 'sankeshwar' in loc_str.lower():
            continue
            
        ip_str = str(local_ip).strip() if local_ip and str(local_ip).strip() != 'NA' else ''
        try:
            port_int = int(float(rtsp_port)) if rtsp_port and str(rtsp_port).strip() != 'NA' else 554
        except ValueError:
            port_int = 554

        try:
            total_ch_int = int(float(total_ch)) if total_ch else 16
        except ValueError:
            total_ch_int = 16

        try:
            used_ch_int = int(float(used_ch)) if used_ch else 0
        except ValueError:
            used_ch_int = 0

        nvrs[loc_str] = {
            'name': loc_str,
            'ip': ip_str,
            'port': port_int,
            'total_channels': total_ch_int,
            'used_channels': used_ch_int,
            'make': str(make or ''),
            'model': str(model or '')
        }

        await nvr_repo.upsert(
            name=loc_str,
            ip_address=ip_str,
            port=port_int,
            total_channels=total_ch_int,
            used_channels=used_ch_int,
            make=str(make or ''),
            model=str(model or '')
        )

    print(f"Loaded {len(nvrs)} valid NVRs from NVR_Factory sheet:")
    for name, n in nvrs.items():
        print(f"  - {name} -> {n['ip']}:{n['port']}")

    # 2. Parse Cameras
    current_dvr_raw = ""
    imported_count = 0
    skipped_count = 0

    for r in range(2, ws_cam.max_row + 1):
        dvr_cell = ws_cam.cell(r, 3).value
        ch_num = ws_cam.cell(r, 5).value
        loc_cell = ws_cam.cell(r, 6).value
        status_cell = ws_cam.cell(r, 7).value

        if dvr_cell and str(dvr_cell).strip():
            current_dvr_raw = str(dvr_cell).strip()

        # Skip Sankeshwar Godowns
        if 'sankeshwar' in current_dvr_raw.lower():
            continue

        if not ch_num or not loc_cell:
            continue

        loc_str = str(loc_cell).strip()
        status_str = str(status_cell).strip() if status_cell else "Working"
        
        # If explicitly marked as 'No Camera' or similar, skip or mark disabled
        is_active = status_str.lower() != 'no camera'

        # Match NVR
        canonical_nvr_name = NVR_ALIASES.get(current_dvr_raw.lower(), current_dvr_raw)
        nvr_info = nvrs.get(canonical_nvr_name)
        
        if not nvr_info:
            # Try fuzzy match
            for k in nvrs.keys():
                if k.lower() in current_dvr_raw.lower() or current_dvr_raw.lower() in k.lower():
                    nvr_info = nvrs[k]
                    canonical_nvr_name = k
                    break

        if not nvr_info:
            print(f"Row {r}: Could not map DVR '{current_dvr_raw}' - skipping")
            skipped_count += 1
            continue

        try:
            ch_int = int(float(ch_num))
        except ValueError:
            ch_int = 1

        ch_formatted = f"{ch_int:02d}"
        cam_name = loc_str
        
        # Build local RTSP URL using Dahua/CP-Plus standard API
        rtsp_url = f"rtsp://arechs_cctv:scpl@2026@{nvr_info['ip']}:{nvr_info['port']}/cam/realmonitor?channel={ch_int}&subtype=0"
        is_no_cam = ('no cam' in loc_str.lower() or 'no camera' in status_str.lower() or not is_active)

        cam_id = await repo.create(
            name=cam_name,
            dvr_nvr_name=canonical_nvr_name,
            location=loc_str,
            ip_address=nvr_info['ip'],
            port=nvr_info['port'],
            channel_no=ch_formatted,
            rtsp_url=rtsp_url,
            is_enabled=is_active,
            is_no_cam=is_no_cam
        )
        if is_no_cam:
            await repo.update_status(cam_id, status="NO_CAM", consecutive_failures=0)
        imported_count += 1

    print(f"\n[DONE] Successfully imported {imported_count} real cameras into database!")
    if skipped_count > 0:
        print(f"Skipped {skipped_count} entries.")

if __name__ == "__main__":
    asyncio.run(import_real_cameras())
