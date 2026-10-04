"""
Automated NVR Channel Sync Script (Enhanced)
Queries both ChannelTitle and RemoteDevice configurations from Dahua / CP Plus NVR firmware.
Any channel with no remote IP (Address == 0.0.0.0 / 192.168.0.0) or named 'NO CAM' is accurately
classified as an unassigned/spare slot (is_no_cam = 1, status = 'NO_CAM').
"""
import shutil
import sqlite3
import httpx
import re
import openpyxl
from datetime import datetime

DB_PATH = "cctv_monitor.db"
EXCEL_PATH = "SCPL_FW_NVR.xlsx"
PASSWORD = "scpl@2026"

def load_nvrs_from_excel():
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb['NVR_Factory']
    nvrs = []
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, 3).value
        ip = ws.cell(r, 7).value
        http_port = ws.cell(r, 11).value
        rtsp_port = ws.cell(r, 12).value
        user = ws.cell(r, 18).value
        if name and ip and str(ip).strip() != 'NA':
            nvrs.append({
                'name': str(name).strip(),
                'ip': str(ip).strip(),
                'http_port': int(float(http_port)) if http_port else 80,
                'rtsp_port': int(float(rtsp_port)) if rtsp_port else 554,
                'user': str(user).strip() if user else 'arechs_cctv'
            })
    return nvrs

def fetch_nvr_config(nvr):
    auth = httpx.DigestAuth(nvr['user'], PASSWORD)
    titles = {}
    remote_ips = {}
    
    with httpx.Client(auth=auth, timeout=5.0) as client:
        # 1. Fetch channel titles
        try:
            url_titles = f"http://{nvr['ip']}:{nvr['http_port']}/cgi-bin/configManager.cgi?action=getConfig&name=ChannelTitle"
            resp = client.get(url_titles)
            if resp.status_code == 200:
                for line in resp.text.splitlines():
                    m = re.match(r'table\.ChannelTitle\[(\d+)\]\.Name=(.*)', line)
                    if m:
                        idx = int(m.group(1))
                        titles[idx + 1] = m.group(2).strip()
        except Exception as e:
            print(f"  [!] Failed fetching ChannelTitle from {nvr['name']}: {e}")

        # 2. Fetch RemoteDevice IP assignments (for NVRs with IP cameras)
        try:
            url_remote = f"http://{nvr['ip']}:{nvr['http_port']}/cgi-bin/configManager.cgi?action=getConfig&name=RemoteDevice"
            resp_rem = client.get(url_remote)
            if resp_rem.status_code == 200:
                for line in resp_rem.text.splitlines():
                    m = re.match(r'table\.RemoteDevice\.uuid:System_CONFIG_NETCAMERA_INFO_(\d+)\.Address=(.*)', line)
                    if m:
                        idx = int(m.group(1))
                        remote_ips[idx + 1] = m.group(2).strip()
        except Exception:
            pass

    return titles, remote_ips

def is_channel_empty(name: str, remote_ip: str = "") -> bool:
    if not name:
        return True
    clean = name.strip().upper()
    if clean in ("NO CAM", "NO CAMERA", "NOCAM", "N/A", "NA", "CHANNEL", "CAM"):
        return True
    if re.match(r'^(D\d+C\d+-)?NO\s*CAM$', clean):
        return True
    if re.match(r'^CHANNEL\d+$', clean):
        return True
    # If the NVR has a dummy IP assigned to this channel (unassigned)
    if remote_ip in ("192.168.0.0", "0.0.0.0"):
        return True
    return False

def sync_fleet():
    nvrs = load_nvrs_from_excel()
    print(f"[*] Found {len(nvrs)} NVR recorders.")

    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()

    total_updated = 0
    total_no_cam_flagged = 0

    for nvr in nvrs:
        print(f"[*] Querying NVR: {nvr['name']} ({nvr['ip']}:{nvr['http_port']})...")
        titles, remote_ips = fetch_nvr_config(nvr)
        if not titles:
            print(f"  [!] Skipped {nvr['name']}: No response from NVR firmware API.")
            continue

        print(f"  [+] Received {len(titles)} titles, {len(remote_ips)} remote device configs.")

        for ch_num, nvr_raw_name in titles.items():
            remote_ip = remote_ips.get(ch_num, "")
            empty = is_channel_empty(nvr_raw_name, remote_ip)
            
            if empty:
                clean_name = "No Cam"
                cursor.execute("""
                    UPDATE cameras 
                    SET name = ?, is_no_cam = 1, status = 'NO_CAM', latency_ms = 0.0, last_error = NULL, consecutive_failures = 0
                    WHERE dvr_nvr_name = ? AND CAST(channel_no AS INTEGER) = ?
                """, (clean_name, nvr['name'], ch_num))
                total_no_cam_flagged += 1
            else:
                clean_name = nvr_raw_name
                cursor.execute("""
                    UPDATE cameras 
                    SET name = ?, is_no_cam = 0
                    WHERE dvr_nvr_name = ? AND CAST(channel_no AS INTEGER) = ?
                """, (clean_name, nvr['name'], ch_num))
            
            total_updated += 1

    # Resolve any open incidents for cameras that are now NO_CAM
    cursor.execute("""
        UPDATE incidents 
        SET resolved_at = datetime('now', 'localtime'), 
            error_reason = 'Channel re-classified as unassigned/spare slot by NVR firmware sync'
        WHERE resolved_at IS NULL AND camera_id IN (
            SELECT id FROM cameras WHERE is_no_cam = 1
        )
    """)
    resolved_incidents = cursor.rowcount
    print(f"\n[*] Cleaned up {resolved_incidents} stale incidents for unassigned slots.")

    conn.commit()
    conn.close()

    print(f"\n[SUCCESS] Enhanced Sync completed!")
    print(f"  Total channels updated: {total_updated}")
    print(f"  Spare slots accurately classified: {total_no_cam_flagged}")

if __name__ == "__main__":
    sync_fleet()
